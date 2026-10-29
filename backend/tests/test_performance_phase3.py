"""
PDFBolt Performance Optimization Suite - Phase 3
Tests for cancellation handling, bounded retries, concurrent office & OCR workloads,
job lifecycle security, and temporary resource cleanup.
"""

import io
import time
import asyncio
from pathlib import Path
from unittest.mock import patch, MagicMock
import pytest
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
import docx
from pptx import Presentation

from backend.app.processors.base import BaseProcessor
from backend.app.processors.ocr import OcrProcessor
from backend.app.processors.word_to_pdf import WordToPdfProcessor
from backend.app.processors.ppt_to_pdf import PptToPdfProcessor
from backend.app.processors.merge import MergeProcessor
from backend.app.services.job_service import job_service
from backend.app.services.queue_service import LocalQueueService
from backend.app.services.cleanup_service import cleanup_service
from backend.app.core.errors import PDFBoltError
from backend.app.core.validation import validate_pdf_output

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def _create_sample_docx(path: Path):
    d = docx.Document()
    d.add_heading("Phase 3 Test Document", level=1)
    d.add_paragraph("Testing concurrent Word to PDF conversion pipeline.")
    d.save(str(path))


def _create_sample_pptx(path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Phase 3 PPT Test"
    prs.save(str(path))


def _create_synthetic_pdf(path: Path, pages: int = 5):
    buf = io.BytesIO()
    can = canvas.Canvas(buf, pagesize=(612, 792))
    for i in range(pages):
        can.drawString(100, 700, f"Page {i+1} synthetic content")
        can.showPage()
    can.save()
    buf.seek(0)
    path.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------
# Test 1: Concurrent Word Conversions with Isolated Profiles
# ---------------------------------------------------------------------------
def test_concurrent_word_conversions_phase3(tmp_path):
    """VERIFICATION: Concurrent Word conversions run in parallel with zero profile collision."""
    from concurrent.futures import ThreadPoolExecutor

    doc_files = []
    for i in range(3):
        p = tmp_path / f"test_word_{i}.docx"
        _create_sample_docx(p)
        doc_files.append(p)

    def _convert(p, idx):
        proc = WordToPdfProcessor(job_id=f"p3_word_{idx}", work_dir=tmp_path / f"w_work_{idx}")
        return proc.run([p], {})

    with ThreadPoolExecutor(max_workers=3) as executor:
        results = list(executor.map(lambda item: _convert(item[1], item[0]), enumerate(doc_files)))

    assert len(results) == 3
    for r in results:
        assert r.status == "COMPLETED"
        assert validate_pdf_output(r.output_path) >= 1


# ---------------------------------------------------------------------------
# Test 2: Concurrent PowerPoint Conversions with Isolated Profiles
# ---------------------------------------------------------------------------
def test_concurrent_ppt_conversions_phase3(tmp_path):
    """VERIFICATION: Concurrent PPT conversions run in parallel without collisions."""
    from concurrent.futures import ThreadPoolExecutor

    ppt_files = []
    for i in range(3):
        p = tmp_path / f"test_ppt_{i}.pptx"
        _create_sample_pptx(p)
        ppt_files.append(p)

    def _convert(p, idx):
        proc = PptToPdfProcessor(job_id=f"p3_ppt_{idx}", work_dir=tmp_path / f"p_work_{idx}")
        return proc.run([p], {})

    with ThreadPoolExecutor(max_workers=3) as executor:
        results = list(executor.map(lambda item: _convert(item[1], item[0]), enumerate(ppt_files)))

    assert len(results) == 3
    for r in results:
        assert r.status == "COMPLETED"
        assert validate_pdf_output(r.output_path) >= 1


# ---------------------------------------------------------------------------
# Test 3: Mixed Workload (Office + PDF + OCR)
# ---------------------------------------------------------------------------
def test_mixed_workload_pipeline_phase3(tmp_path):
    """VERIFICATION: Mixed concurrent workloads (Merge + Word + OCR) execute without deadlock."""
    from concurrent.futures import ThreadPoolExecutor

    docx_p = tmp_path / "mix.docx"
    _create_sample_docx(docx_p)

    pdf_p = tmp_path / "mix.pdf"
    _create_synthetic_pdf(pdf_p, pages=3)

    def _run_task(task_type):
        if task_type == "word":
            p = WordToPdfProcessor(job_id="mix_word", work_dir=tmp_path / "w_mix")
            return p.run([docx_p], {})
        elif task_type == "merge":
            p = MergeProcessor(job_id="mix_merge", work_dir=tmp_path / "m_mix")
            return p.run([pdf_p, pdf_p], {})
        else:
            with patch.object(OcrProcessor, "_ocr_image_to_text", return_value="Mixed text"):
                p = OcrProcessor(job_id="mix_ocr", work_dir=tmp_path / "o_mix")
                return p.run([pdf_p], {})

    with ThreadPoolExecutor(max_workers=3) as executor:
        results = list(executor.map(_run_task, ["word", "merge", "ocr"]))

    assert len(results) == 3
    for r in results:
        assert r.status == "COMPLETED"
        assert validate_pdf_output(r.output_path) >= 1


# ---------------------------------------------------------------------------
# Test 4: Job Cancellation & Early Loop Abort
# ---------------------------------------------------------------------------
def test_processor_cancellation_token(tmp_path):
    """VERIFICATION: Setting is_cancelled aborts processor loop early and stops work."""
    pdf_p = tmp_path / "cancel_test.pdf"
    _create_synthetic_pdf(pdf_p, pages=10)

    processor = OcrProcessor(job_id="test_cancel", work_dir=tmp_path / "cancel_work")

    # Cancel after 2 pages
    page_count = 0
    def mock_ocr(img, lang="eng", timeout=30):
        nonlocal page_count
        page_count += 1
        if page_count >= 2:
            processor.cancel()
        return f"Page {page_count} text"

    with patch.object(OcrProcessor, "_ocr_image_to_text", side_effect=mock_ocr):
        with pytest.raises(PDFBoltError) as exc_info:
            processor.process([pdf_p], {"force_ocr": True})

    assert exc_info.value.code == "JOB_CANCELLED"
    assert processor.is_cancelled.is_set()
    assert page_count < 10, f"Expected early cancellation, but processed {page_count} pages"


# ---------------------------------------------------------------------------
# Test 5: Queue Service Bounded Retries (Poison Pill Prevention)
# ---------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_queue_service_bounded_retries():
    """VERIFICATION: Poison-pill jobs exceeding max retries move to dead-letter queue."""
    queue = LocalQueueService()
    job_id = await queue.enqueue({"job_id": "test_retry_job", "operation": "merge"})

    # Fetch and retry 3 times
    for _ in range(3):
        j = await queue.get()
        assert j is not None
        await queue.retry(j["job_id"])

    # 4th retry should trigger dead-letter
    j = await queue.get()
    assert j is not None
    await queue.retry(j["job_id"])

    # Queue should now be empty (moved to dead letter)
    next_j = await queue.get()
    assert next_j is None


# ---------------------------------------------------------------------------
# Test 6: Job Token Authorization & User Privacy Isolation
# ---------------------------------------------------------------------------
def test_job_token_isolation():
    """VERIFICATION: Valid download token is strictly required to access job results."""
    from backend.app.api.v1.jobs import _require_job_token
    from fastapi import HTTPException
    from backend.app.models.job import Job, JobStatus
    import secrets

    job_id = "test_auth_job"
    token = secrets.token_urlsafe(32)
    fake_job = Job(
        job_id=job_id,
        status=JobStatus.COMPLETED,
        operation="merge",
        download_token=token
    )
    job_service.jobs[job_id] = fake_job

    # Correct token -> No exception
    _require_job_token(job_id, token)

    # Wrong token -> 401 Unauthorized
    with pytest.raises(HTTPException) as exc:
        _require_job_token(job_id, "wrong_token_12345")
    assert exc.value.status_code == 401

    # Cleanup
    job_service.jobs.pop(job_id, None)


# ---------------------------------------------------------------------------
# Test 7: Safe Cleanup Isolation (Cannot Delete Another Job's Files)
# ---------------------------------------------------------------------------
def test_cleanup_service_isolation(tmp_path):
    """VERIFICATION: delete_job_files only deletes its own job directory and preserves peers."""
    custom_cleanup = cleanup_service.__class__(storage_dir=str(tmp_path))

    job1_dir = tmp_path / "job_alpha"
    job1_dir.mkdir()
    (job1_dir / "output.pdf").write_bytes(b"Job 1 PDF Data")

    job2_dir = tmp_path / "job_beta"
    job2_dir.mkdir()
    (job2_dir / "output.pdf").write_bytes(b"Job 2 PDF Data")

    # Delete job_alpha
    custom_cleanup.delete_job_files("job_alpha")

    # Verify job_alpha was deleted, but job_beta remains untouched
    assert not job1_dir.exists()
    assert job2_dir.exists()
    assert (job2_dir / "output.pdf").read_bytes() == b"Job 2 PDF Data"
